#!/usr/bin/env python3
"""
ScamShield AI - Reference Flask Backend Service
Integrates with sample.xlsx company master database and dataset.json threat feed.
Triggers high-risk warnings whenever an INACTIVE domain/company is opened.
"""

import json
import os
import sys
import urllib.parse
from flask import Flask, request, jsonify, render_template_string, Response

app = Flask(__name__)

# File Paths
BACKEND_DIR = os.path.dirname(os.path.abspath(__file__))
DATASET_JSON_FILE = os.path.join(BACKEND_DIR, "dataset.json")

def find_sample_xlsx():
    candidate_paths = [
        os.path.join(BACKEND_DIR, "sample.xlsx"),
        os.path.join(BACKEND_DIR, "..", "sample.xlsx"),
        "d:/extention/sample.xlsx",
        "d:/extention/backend/sample.xlsx"
    ]
    for p in candidate_paths:
        if os.path.exists(p):
            return os.path.abspath(p)
    return None

# Cache for sample.xlsx
_xlsx_cache = {
    "mtime": 0,
    "records": []
}

def load_dataset_json():
    if os.path.exists(DATASET_JSON_FILE):
        try:
            with open(DATASET_JSON_FILE, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception as e:
            print(f"Error reading dataset.json: {e}", file=sys.stderr)
    return []

def load_sample_xlsx():
    xlsx_path = find_sample_xlsx()
    if not xlsx_path:
        print("Warning: sample.xlsx not found on disk", file=sys.stderr)
        return []

    try:
        current_mtime = os.path.getmtime(xlsx_path)
        if current_mtime == _xlsx_cache["mtime"] and _xlsx_cache["records"]:
            return _xlsx_cache["records"]

        import openpyxl
        wb = openpyxl.load_workbook(xlsx_path, data_only=True)
        sheet = wb["Company_Master_Sample"] if "Company_Master_Sample" in wb.sheetnames else wb.active

        records = []
        for r in range(2, sheet.max_row + 1):
            cin = str(sheet.cell(r, 1).value or "").strip()
            name = str(sheet.cell(r, 2).value or "").strip()
            status = str(sheet.cell(r, 3).value or "").strip()
            brand = str(sheet.cell(r, 9).value or "").strip()
            official_domain = str(sheet.cell(r, 10).value or "").strip()
            official_website = str(sheet.cell(r, 11).value or "").strip()

            # Ignore empty / invalid rows
            if not status or not name or status.startswith("#") or name.startswith("#"):
                continue

            # Extract normalized domains from official domain and website
            domains = set()
            for raw in [official_domain, official_website]:
                if raw and not raw.startswith("#"):
                    d = normalize_domain_backend(raw)
                    if d:
                        domains.add(d)

            records.append({
                "cin": cin,
                "name": name,
                "status": status,
                "brand": brand,
                "official_domain": official_domain,
                "official_website": official_website,
                "domains": list(domains)
            })

        _xlsx_cache["mtime"] = current_mtime
        _xlsx_cache["records"] = records
        print(f"Loaded {len(records)} records from {xlsx_path}", file=sys.stderr)
        return records

    except Exception as e:
        print(f"Error loading sample.xlsx: {e}", file=sys.stderr)
        return _xlsx_cache["records"]

def normalize_domain_backend(raw_domain, raw_url=None):
    """
    Independently validate and normalize domain on backend.
    Strips schemes, paths, query strings, ports, and leading 'www.'.
    """
    domain = (raw_domain or "").strip().lower()

    if raw_url:
        try:
            parsed = urllib.parse.urlparse(raw_url.strip())
            if parsed.hostname:
                domain = parsed.hostname.lower()
        except Exception:
            pass

    if "://" in domain:
        domain = domain.split("://", 1)[1]
    if "/" in domain:
        domain = domain.split("/", 1)[0]
    if ":" in domain:
        domain = domain.split(":", 1)[0]
    if domain.startswith("www."):
        domain = domain[4:]

    return domain.strip()

# Global CORS Handler
@app.after_request
def add_cors_headers(response: Response):
    response.headers["Access-Control-Allow-Origin"] = "*"
    response.headers["Access-Control-Allow-Methods"] = "GET, POST, OPTIONS"
    response.headers["Access-Control-Allow-Headers"] = "Content-Type, Authorization, Accept"
    return response

@app.route("/api/check-url", methods=["OPTIONS"])
def check_url_options():
    return Response(status=204)

@app.route("/api/check-url", methods=["POST"])
def check_url():
    """
    Check submitted domain/URL against:
    1. sample.xlsx dataset (Checking Company Status: Inactive vs Active)
    2. dataset.json threat feed
    """
    data = request.get_json(silent=True) or {}
    raw_url = data.get("url", "")
    raw_domain = data.get("domain", "")

    domain = normalize_domain_backend(raw_domain, raw_url)

    if not domain:
        return jsonify({
            "found": False,
            "risk_score": 0,
            "risk_level": "UNKNOWN",
            "reason": "Invalid or empty domain provided",
            "domain": ""
        }), 400

    # 1. Check against sample.xlsx
    xlsx_records = load_sample_xlsx()
    for rec in xlsx_records:
        matched = False
        for d in rec["domains"]:
            if domain == d or domain.endswith("." + d) or d.endswith("." + domain):
                matched = True
                break

        if matched:
            status = rec["status"].strip()
            # If status is Inactive -> Trigger WARNING
            if status.lower() in ["inactive", "dormant", "struck off", "dissolved", "strike off", "under liquidation", "closed"]:
                return jsonify({
                    "found": True,
                    "risk_score": 90,
                    "risk_level": "HIGH",
                    "status": "Inactive",
                    "company_status": "Inactive",
                    "company_name": rec["name"],
                    "cin": rec["cin"],
                    "source": "sample.xlsx (Company Master Dataset)",
                    "reason": f"Company '{rec['name']}' is marked as INACTIVE in dataset (sample.xlsx). Caution: Website or company may be defunct or unauthorized.",
                    "domain": domain
                }), 200
            else:
                # Active in master dataset -> Verified Active, No Threat Match
                return jsonify({
                    "found": False,
                    "verified_active": True,
                    "status": "Active",
                    "company_status": "Active",
                    "company_name": rec["name"],
                    "cin": rec["cin"],
                    "risk_level": "LOW",
                    "risk_score": 0,
                    "source": "sample.xlsx (Company Master Dataset)",
                    "reason": f"Company '{rec['name']}' is verified as ACTIVE in dataset (sample.xlsx).",
                    "domain": domain
                }), 200

    # 2. Check against dataset.json (Known threat domains)
    threat_dataset = load_dataset_json()
    for entry in threat_dataset:
        entry_domain = normalize_domain_backend(entry.get("domain", ""), entry.get("url", ""))
        if entry_domain == domain or domain.endswith("." + entry_domain):
            return jsonify({
                "found": True,
                "risk_score": entry.get("risk_score", 85),
                "risk_level": entry.get("risk_level", "HIGH"),
                "reason": entry.get("reason", "Domain found in ScamShield threat dataset"),
                "domain": domain
            }), 200

    # 3. No match found in either database
    return jsonify({
        "found": False,
        "risk_score": 0,
        "risk_level": "UNKNOWN",
        "reason": "No known match found in ScamShield dataset",
        "domain": domain
    }), 200

# ===================================================
# MOCK SCAMSHIELD WEB APPLICATION PAGES
# ===================================================

HTML_HOME_TEMPLATE = """<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8">
  <title>ScamShield AI — Intelligence Platform</title>
  <style>
    body {
      margin: 0;
      background: #00122e;
      color: #f8fafc;
      font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif;
      display: flex;
      flex-direction: column;
      align-items: center;
      justify-content: center;
      min-height: 100vh;
      padding: 20px;
    }
    .card {
      background: #04193d;
      border: 1px solid #233b63;
      padding: 30px 40px;
      max-width: 600px;
      box-shadow: 0 10px 30px rgba(0,0,0,0.5);
      text-align: center;
    }
    h1 {
      color: #38bdf8;
      margin-top: 0;
    }
    p {
      color: #9bb1cf;
      line-height: 1.6;
    }
  </style>
</head>
<body>
  <div class="card">
    <h1>ScamShield AI Intelligence Platform</h1>
    <p>The ScamShield AI browser extension is successfully connected to the dataset and backend server.</p>
  </div>
</body>
</html>"""

HTML_ANALYZE_TEMPLATE = """<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8">
  <title>ScamShield AI — Threat Analysis</title>
  <style>
    body {
      margin: 0;
      background: #00122e;
      color: #f8fafc;
      font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif;
      display: flex;
      flex-direction: column;
      align-items: center;
      justify-content: center;
      min-height: 100vh;
      padding: 20px;
    }
    .card {
      background: #4a0d11;
      border: 1px solid #dc2626;
      padding: 30px 40px;
      max-width: 650px;
      box-shadow: 0 10px 30px rgba(220, 38, 38, 0.3);
    }
    h1 {
      color: #ef4444;
      margin-top: 0;
    }
    .url-box {
      background: #200507;
      border: 1px solid #7f1d1d;
      padding: 12px;
      font-family: monospace;
      color: #fca5a5;
      word-break: break-all;
      margin: 15px 0;
    }
    p {
      color: #cbd5e1;
      line-height: 1.6;
    }
  </style>
</head>
<body>
  <div class="card">
    <h1>⚠ ScamShield Threat & Dataset Report</h1>
    <p>Target domain / URL:</p>
    <div class="url-box">{{ target_url }}</div>
    <p>This website was flagged as INACTIVE or suspicious in the dataset. Proceed with caution.</p>
  </div>
</body>
</html>"""

@app.route("/", methods=["GET"])
def home():
    return render_template_string(HTML_HOME_TEMPLATE)

@app.route("/analyze", methods=["GET"])
def analyze():
    target_url = request.args.get("url", "No URL provided")
    return render_template_string(HTML_ANALYZE_TEMPLATE, target_url=target_url)

if __name__ == "__main__":
    port = 5000
    if len(sys.argv) > 1 and sys.argv[1].isdigit():
        port = int(sys.argv[1])
    print(f"Starting ScamShield Flask Backend on http://127.0.0.1:{port}")
    app.run(host="127.0.0.1", port=port, debug=False)
