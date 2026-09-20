import os
import openpyxl
import urllib.parse

def load_sample_xlsx():
    paths = [
        "d:/extention/sample.xlsx",
        "d:/extention/backend/sample.xlsx",
        os.path.join(os.path.dirname(__file__), "sample.xlsx"),
        os.path.join(os.path.dirname(__file__), "..", "sample.xlsx")
    ]
    file_path = None
    for p in paths:
        if os.path.exists(p):
            file_path = p
            break
    
    if not file_path:
        print("sample.xlsx not found!")
        return []

    wb = openpyxl.load_workbook(file_path, data_only=True)
    sheet = wb["Company_Master_Sample"] if "Company_Master_Sample" in wb.sheetnames else wb.active
    
    records = []
    for r in range(2, sheet.max_row + 1):
        cin = str(sheet.cell(r, 1).value or "").strip()
        name = str(sheet.cell(r, 2).value or "").strip()
        status = str(sheet.cell(r, 3).value or "").strip()
        brand = str(sheet.cell(r, 9).value or "").strip()
        official_domain = str(sheet.cell(r, 10).value or "").strip()
        official_website = str(sheet.cell(r, 11).value or "").strip()
        
        if not name and not cin and not official_website:
            continue

        domains = set()
        for raw in [official_domain, official_website]:
            if raw:
                # normalize
                d = raw.lower().strip()
                if "://" in d:
                    try:
                        parsed = urllib.parse.urlparse(d)
                        if parsed.hostname:
                            d = parsed.hostname.lower()
                    except Exception:
                        pass
                if "/" in d:
                    d = d.split("/")[0]
                if ":" in d:
                    d = d.split(":")[0]
                if d.startswith("www."):
                    d = d[4:]
                if d:
                    domains.add(d)

        records.append({
            "cin": cin,
            "name": name,
            "status": status,
            "brand": brand,
            "official_domain": official_domain,
            "official_website": official_website,
            "domains": list(domains)
        })

    return records

if __name__ == "__main__":
    records = load_sample_xlsx()
    print(f"Loaded {len(records)} records from sample.xlsx:")
    for rec in records:
        print(f"Status: {rec['status']:<10} | Domains: {rec['domains']} | Name: {rec['name']}")
