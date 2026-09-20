import openpyxl
import urllib.parse

import os

excel_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "sample.xlsx")
wb = openpyxl.load_workbook(excel_path, data_only=True)
sheet = wb["Company_Master_Sample"]

print("Columns:")
for c in range(1, sheet.max_column + 1):
    print(c, sheet.cell(1, c).value)

print("\nRows:")
for r in range(2, sheet.max_row + 1):
    cin = sheet.cell(r, 1).value
    name = sheet.cell(r, 2).value
    status = sheet.cell(r, 3).value
    domain = sheet.cell(r, 10).value
    url = sheet.cell(r, 11).value
    if name or status or cin:
        print(f"Row {r:2d} | CIN: {cin} | Status: {status} | Domain: {domain} | URL: {url} | Name: {name}")
